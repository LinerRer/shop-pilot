import asyncio
import sys
import json
import datetime
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parents[1]))

from docx import Document
from docx.shared import Pt

from utils.config import settings
from analytics import metrics
from langchain_deepseek import ChatDeepSeek
from langchain_mcp_adapters.client import MultiServerMCPClient
from langgraph.prebuilt import create_react_agent
from langchain_core.messages import SystemMessage, HumanMessage

SERVER = Path(__file__).resolve().parents[1] / "mcp_servers" / "ecomdb_server.py"
OUT_DIR = Path(__file__).resolve().parents[2] / "output"

HEALTH_SYSTEM = """你是一名资深电商运营分析师。基于确定性异常清单做根因归因，输出中文报告：
1. 对清单中 severity=high 的异常商品，调用 get_product_detail(product_id, days=14) 查看库存和差评原文。
2. 归因规则：库存0已售罄=断货补货；差评集中破损/漏水/快递=物流包装问题；差评集中质量/尺码=商品与详情页问题。
3. 不要重复列异常清单（清单已由系统生成表格），只写：根因判断 + 每条可执行建议。
4. 所有结论基于工具返回，禁止编造数字。"""


async def get_root_cause_text():
    client = MultiServerMCPClient({
        "ecomdb": {"command": sys.executable, "args": [str(SERVER)], "transport": "stdio"}
    })
    tools = await client.get_tools()
    llm = ChatDeepSeek(
        api_key=settings.deepseek_api_key,
        base_url=settings.deepseek_base_url,
        model=settings.deepseek_model,
        temperature=0,
    )
    agent = create_react_agent(llm, tools)
    resp = await agent.ainvoke({"messages": [
        SystemMessage(content=HEALTH_SYSTEM),
        HumanMessage(content="先用 scan_anomalies 拿到异常清单，然后对 high 级商品做根因分析并给出建议。"),
    ]})
    return resp["messages"][-1].content


def build_docx(anomalies, ai_text, out_path):
    doc = Document()
    doc.add_heading("电商智能运营日报", level=0)
    doc.add_paragraph(f"生成日期：{datetime.date.today().isoformat()}")

    doc.add_heading("一、异常清单（近7天）", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "对象"
    hdr[1].text = "异常类型"
    hdr[2].text = "数据证据"
    hdr[3].text = "级别"
    for a in anomalies:
        row = table.add_row().cells
        row[0].text = f"{a.get('scope', '')} {a.get('product', '')}".strip()
        row[1].text = a["type"]
        row[2].text = a["evidence"]
        row[3].text = a["severity"]

    doc.add_heading("二、根因分析与建议", level=1)
    for line in ai_text.split("\n"):
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        doc.add_paragraph(line.lstrip("-* "))

    doc.save(out_path)


async def main():
    OUT_DIR.mkdir(exist_ok=True)
    anomalies = metrics.detect_anomalies()
    print(f"确定性扫描到 {len(anomalies)} 条异常，正在让Agent做根因归因...")
    ai_text = await get_root_cause_text()
    out_path = OUT_DIR / f"运营日报_{datetime.date.today().isoformat()}.docx"
    build_docx(anomalies, ai_text, out_path)
    print(f"日报已生成：{out_path}")

    from database.models import engine, DailyReport
    from sqlalchemy.orm import Session
    high_count = sum(1 for a in anomalies if a["severity"] == "high")
    with Session(engine) as sess:
        sess.add(DailyReport(
            report_date=datetime.date.today(),
            anomaly_count=len(anomalies),
            high_count=high_count,
            anomalies_json=json.dumps(anomalies, ensure_ascii=False),
            root_cause=ai_text,
        ))
        sess.commit()
    print(f"已入库：异常{len(anomalies)}条 / high {high_count}条")


if __name__ == "__main__":
    asyncio.run(main())
